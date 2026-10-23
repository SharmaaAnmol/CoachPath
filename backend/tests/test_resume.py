"""
Comprehensive Test Suite for Phase 5A: Resume Upload & AI Parsing Backend.
Validates file validation, MIME magic bytes, document text extraction (PDF & DOCX),
structured AI extraction, skill normalization, human review gate (deferred profile merging),
multi-tenant authorization, retry mechanics, download, and deletion.
"""
from __future__ import annotations

import io
import uuid
import docx
import pytest
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai.factory import override_llm_provider, reset_llm_provider
from app.ai.mock_adapter import MockLLMProvider
from app.models.career import (
    CareerProfile,
    Certification,
    Education,
    Experience,
    Project,
)
from app.models.resume import Resume
from app.models.skills import Skill, UserSkill
from app.schemas.resume import ParsingErrorCategory, ResumeStatus
from app.services.storage import storage_service


# ─── Sample Binary Generators ────────────────────────────────────────

def make_valid_pdf_bytes(
    candidate_name: str = "Aarav Sharma",
    content_text: str = (
        "Senior Backend Engineer with 5 years experience in Python, FastAPI, and PostgreSQL. "
        "Education: Bachelor of Technology in Computer Science from State University. "
        "Projects: Built distributed API rate limiter with Redis and Docker. "
        "Certifications: AWS Certified Solutions Architect."
    ),
) -> bytes:
    """Generates valid PDF binary stream containing readable text."""
    text_escaped = f"{candidate_name}\n{content_text}".replace("(", "\\(").replace(")", "\\)")
    stream_content = f"BT /F1 12 Tf 72 720 Td ({text_escaped}) Tj ET".encode("latin1")
    length = len(stream_content)

    return f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length {length} >>
stream
{stream_content.decode('latin1')}
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000400 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
460
%%EOF""".encode("latin1")


def make_blank_image_pdf_bytes() -> bytes:
    """Generates PDF with no extractable text (simulating scanned/image document)."""
    return b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>
endobj
xref
0 4
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
trailer
<< /Size 4 /Root 1 0 R >>
startxref
190
%%EOF"""


def make_valid_docx_bytes(
    candidate_name: str = "Kavya Patel",
    headline: str = "Lead Cloud Engineer",
) -> bytes:
    """Generates valid DOCX document using python-docx."""
    doc = docx.Document()
    doc.add_heading(candidate_name, 0)
    doc.add_paragraph(headline)
    doc.add_paragraph(
        "Summary: Experienced software engineer specializing in distributed systems, "
        "Python, Docker, Kubernetes, and PostgreSQL database optimization."
    )
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Master of Science in Software Engineering, Tech University (2019-2021).")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Senior Engineer at CloudCorp (2021 - Present). Designed multi-tenant microservices.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def mock_ai_resume_extraction():
    """Generates structured ExtractedResumeSchema dictionary."""
    return {
        "personal_info": {
            "name": "Aarav Sharma",
            "email": "aarav.sharma@example.com",
            "phone": "+1-555-0199",
            "location": "Bangalore, India",
            "linkedin_url": "https://linkedin.com/in/aaravsharma",
            "github_url": "https://github.com/aaravsharma",
            "portfolio_url": None,
        },
        "education": [
            {
                "institution": "State University",
                "degree": "Bachelor of Technology",
                "field_of_study": "Computer Science",
                "start_date": "2016",
                "end_date": "2020",
                "grade_or_gpa": "3.8 / 4.0",
                "highlights": ["Dean's Honor List"],
                "evidence": {
                    "section_name": "Education",
                    "page_number": 1,
                    "supporting_excerpt": "Bachelor of Technology in Computer Science from State University",
                },
            }
        ],
        "experience": [
            {
                "employer": "TechCorp Systems",
                "job_title": "Backend Software Engineer",
                "start_date": "2020-07",
                "end_date": "2023-08",
                "is_current": False,
                "location": "Bangalore",
                "responsibilities": ["Engineered high-throughput REST APIs with FastAPI"],
                "achievements": ["Reduced endpoint latency by 35% through Redis caching"],
                "evidence": {
                    "section_name": "Experience",
                    "page_number": 1,
                    "supporting_excerpt": "Backend Software Engineer at TechCorp Systems",
                },
            }
        ],
        "skills": [
            {"name": "Python", "category": "language", "confidence_source": "resume_stated", "evidence": None},
            {"name": "FastAPI", "category": "framework", "confidence_source": "resume_stated", "evidence": None},
            {"name": "PostgreSQL", "category": "database", "confidence_source": "resume_stated", "evidence": None},
            {"name": "Docker", "category": "tool", "confidence_source": "resume_stated", "evidence": None},
            {"name": "Redis", "category": "database", "confidence_source": "resume_stated", "evidence": None},
            {"name": "QuantumComputingX", "category": "concept", "confidence_source": "resume_stated", "evidence": None},
        ],
        "projects": [
            {
                "name": "API Rate Limiter",
                "description": "Distributed token bucket rate limiter",
                "technologies": ["Python", "Redis", "Docker"],
                "url": "https://github.com/aaravsharma/rate-limiter",
                "highlights": ["Handled 50k requests per second benchmark"],
                "evidence": None,
            }
        ],
        "certifications": [
            {
                "name": "AWS Certified Solutions Architect",
                "issuing_organization": "Amazon Web Services",
                "issue_date": "2022-03",
                "expiration_date": "2025-03",
                "credential_id": "AWS-123456",
                "evidence": None,
            }
        ],
        "executive_summary": "Senior Backend Engineer with 5 years experience in Python, FastAPI, and PostgreSQL.",
        "additional_sections": {},
        "uncertain_items": [],
    }


async def register_and_login(client: AsyncClient, email: str, full_name: str) -> str:
    """Helper creating user account and returning valid JWT bearer token."""
    pw = "SecurePassword123!"
    reg = await client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": pw, "full_name": full_name},
    )
    assert reg.status_code == 201
    login = await client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": pw},
    )
    assert login.status_code == 200
    return login.json()["access_token"]


@pytest.fixture(autouse=True)
def setup_mock_llm():
    """Installs MockLLMProvider for all tests in this module."""
    mock = MockLLMProvider()
    mock.set_structured_response(mock_ai_resume_extraction())
    override_llm_provider(mock)
    yield mock
    reset_llm_provider()


# ─── 1. File Upload Tests ─────────────────────────────────────────────


@pytest.mark.asyncio
async def test_upload_valid_pdf_success(client: AsyncClient):
    """Uploading a valid text-based PDF successfully parses and stores draft."""
    token = await register_and_login(client, "pdf.upload@example.com", "Aarav Sharma")
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = make_valid_pdf_bytes("Aarav Sharma")
    files = {"file": ("aarav_resume.pdf", pdf_bytes, "application/pdf")}

    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 201
    data = r.json()
    assert data["status"] == "parsed"
    assert data["file_name"] == "aarav_resume.pdf"
    assert data["mime_type"] == "application/pdf"
    assert "resume_id" in data


@pytest.mark.asyncio
async def test_upload_valid_docx_success(client: AsyncClient):
    """Uploading a valid DOCX document successfully extracts text and parses."""
    token = await register_and_login(client, "docx.upload@example.com", "Kavya Patel")
    headers = {"Authorization": f"Bearer {token}"}

    docx_bytes = make_valid_docx_bytes("Kavya Patel")
    files = {
        "file": (
            "kavya_resume.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }

    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 201
    data = r.json()
    assert data["status"] == "parsed"
    assert data["file_name"] == "kavya_resume.docx"
    assert "resume_id" in data


@pytest.mark.asyncio
async def test_upload_unsupported_file_extension_rejected(client: AsyncClient):
    """Uploading an unsupported file format (.txt, .exe, .png) is rejected with HTTP 400."""
    token = await register_and_login(client, "badext.upload@example.com", "Bad Ext User")
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("malicious.txt", b"Plain text content", "text/plain")}
    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 400
    assert "Only PDF (.pdf) and Word (.docx) files are supported" in r.json()["error"]["message"]


@pytest.mark.asyncio
async def test_upload_spoofed_mime_type_rejected(client: AsyncClient):
    """Uploading a text or executable file renamed as .pdf is rejected by magic byte validation."""
    token = await register_and_login(client, "spoofed.upload@example.com", "Spoofed User")
    headers = {"Authorization": f"Bearer {token}"}

    # Text content disguised as .pdf
    files = {"file": ("fake_resume.pdf", b"Hello, this is just plain text, not a PDF!", "application/pdf")}
    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 400
    assert "binary signature" in r.json()["error"]["message"]


@pytest.mark.asyncio
async def test_upload_empty_file_rejected(client: AsyncClient):
    """Uploading a zero-byte file is rejected with HTTP 400."""
    token = await register_and_login(client, "empty.upload@example.com", "Empty File User")
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("empty.pdf", b"", "application/pdf")}
    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 400
    assert "empty" in r.json()["error"]["message"].lower()


@pytest.mark.asyncio
async def test_upload_file_exceeding_size_limit_rejected(client: AsyncClient):
    """Uploading a file exceeding 10MB limit is rejected with HTTP 400."""
    token = await register_and_login(client, "oversize.upload@example.com", "Oversize User")
    headers = {"Authorization": f"Bearer {token}"}

    # Generate 10MB + 100 bytes
    huge_bytes = b"%PDF-" + b"0" * (10 * 1024 * 1024 + 100)
    files = {"file": ("huge_resume.pdf", huge_bytes, "application/pdf")}
    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 400
    assert "exceeds maximum allowable limit" in r.json()["error"]["message"]


@pytest.mark.asyncio
async def test_upload_image_only_pdf_flags_ocr_required(client: AsyncClient):
    """A blank or scanned image-only PDF is saved with status='failed' and category='INSUFFICIENT_TEXT_OCR_REQUIRED'."""
    token = await register_and_login(client, "scanned.upload@example.com", "Scanned User")
    headers = {"Authorization": f"Bearer {token}"}

    blank_pdf = make_blank_image_pdf_bytes()
    files = {"file": ("scanned_resume.pdf", blank_pdf, "application/pdf")}

    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 201
    data = r.json()
    assert data["status"] == "failed"

    # Verify detail status
    detail_res = await client.get(f"/api/v1/resume/{data['resume_id']}", headers=headers)
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert detail["status"] == "failed"
    assert detail["parsing_error_category"] == ParsingErrorCategory.INSUFFICIENT_TEXT_OCR_REQUIRED


@pytest.mark.asyncio
async def test_upload_corrupt_pdf_file_flags_corrupted(client: AsyncClient):
    """A malformed/corrupted PDF is handled safely with status='failed' and category='CORRUPTED_DOCUMENT'."""
    token = await register_and_login(client, "corrupt.upload@example.com", "Corrupt User")
    headers = {"Authorization": f"Bearer {token}"}

    # Starts with %PDF- but is corrupted trash
    corrupt_bytes = b"%PDF-1.4\nGARBAGE_NOT_A_VALID_PDF_STRUCTURE"
    files = {"file": ("corrupt.pdf", corrupt_bytes, "application/pdf")}

    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 201
    data = r.json()
    assert data["status"] == "failed"

    detail_res = await client.get(f"/api/v1/resume/{data['resume_id']}", headers=headers)
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert detail["status"] == "failed"
    assert detail["parsing_error_category"] == ParsingErrorCategory.CORRUPTED_DOCUMENT


# ─── 2. AI Parsing & Provider Resilience Tests ───────────────────────


@pytest.mark.asyncio
async def test_upload_ai_provider_failure_flags_service_unavailable(client: AsyncClient):
    """When the AI provider fails (503), the document status is recorded as failed without throwing uncaught 500."""
    from app.ai.openai_adapter import OpenAIProviderAdapter

    unavailable_provider = OpenAIProviderAdapter(api_key="", model="gpt-4o-mini")
    override_llm_provider(unavailable_provider)

    try:
        token = await register_and_login(client, "noai.resume@example.com", "No AI User")
        headers = {"Authorization": f"Bearer {token}"}

        pdf_bytes = make_valid_pdf_bytes("No AI Candidate")
        files = {"file": ("resume_no_ai.pdf", pdf_bytes, "application/pdf")}

        r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
        assert r.status_code == 201
        data = r.json()
        assert data["status"] == "failed"

        detail_res = await client.get(f"/api/v1/resume/{data['resume_id']}", headers=headers)
        assert detail_res.status_code == 200
        detail = detail_res.json()
        assert detail["status"] == "failed"
        assert detail["parsing_error_category"] == ParsingErrorCategory.AI_PROVIDER_UNAVAILABLE
    finally:
        reset_llm_provider()


@pytest.mark.asyncio
async def test_retry_parsing_recovers_failed_resume(client: AsyncClient, setup_mock_llm: MockLLMProvider):
    """Retrying parsing on a failed document re-runs the pipeline and recovers to parsed status."""
    from app.ai.openai_adapter import OpenAIProviderAdapter

    # 1. First upload with unavailable AI -> fails
    unavailable_provider = OpenAIProviderAdapter(api_key="", model="gpt-4o-mini")
    override_llm_provider(unavailable_provider)

    token = await register_and_login(client, "retry.resume@example.com", "Retry Candidate")
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = make_valid_pdf_bytes("Retry Candidate")
    files = {"file": ("retry_resume.pdf", pdf_bytes, "application/pdf")}

    r = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert r.status_code == 201
    resume_id = r.json()["resume_id"]
    assert r.json()["status"] == "failed"

    # 2. Restore available mock AI
    override_llm_provider(setup_mock_llm)

    # 3. Call retry endpoint
    retry_res = await client.post(f"/api/v1/resume/{resume_id}/retry", headers=headers)
    assert retry_res.status_code == 200
    retry_data = retry_res.json()
    assert retry_data["status"] == "parsed"

    # Verify parsed draft is now available
    draft_res = await client.get(f"/api/v1/resume/{resume_id}/parsed", headers=headers)
    assert draft_res.status_code == 200
    draft = draft_res.json()
    assert draft["status"] == "parsed"
    assert draft["extracted_preview"] is not None
    assert len(draft["extracted_preview"]["experience"]) == 1


# ─── 3. Skill Normalization & Human Review Gate Tests ─────────────────


@pytest.mark.asyncio
async def test_skill_normalization_maps_canonical_skills(
    client: AsyncClient,
    test_session: AsyncSession,
):
    """Verifies that extracted skills match canonical skills in the taxonomy database."""
    # Seed a canonical skill into test database if not present
    existing_skill = await test_session.scalar(select(Skill).where(Skill.id == "python"))
    if not existing_skill:
        test_session.add(
            Skill(
                id="python",
                name="Python",
                category="language",
                description="General-purpose high-level programming language",
                synonyms=["python3", "py"],
            )
        )
        test_session.add(
            Skill(
                id="postgresql",
                name="PostgreSQL",
                category="database",
                description="Relational database management system",
                synonyms=["postgres", "pgsql"],
            )
        )
        await test_session.commit()

    token = await register_and_login(client, "normalizer@example.com", "Normalizer User")
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = make_valid_pdf_bytes("Normalizer User")
    files = {"file": ("norm_resume.pdf", pdf_bytes, "application/pdf")}

    upload_res = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["resume_id"]

    draft_res = await client.get(f"/api/v1/resume/{resume_id}/parsed", headers=headers)
    assert draft_res.status_code == 200
    preview = draft_res.json()["extracted_preview"]

    # Python should be normalized
    matched_ids = [ns["skill_id"] for ns in preview["normalized_skills"]]
    assert "python" in matched_ids
    # Unknown skill QuantumComputingX should be in unmatched
    assert "QuantumComputingX" in preview["raw_unmatched_skills"]


@pytest.mark.asyncio
async def test_career_profile_is_not_mutated_on_parsing(
    client: AsyncClient,
    test_session: AsyncSession,
):
    """
    CRITICAL INVARIANT TEST:
    Uploading and parsing a resume must NOT automatically write to, overwrite,
    or mutate the user's canonical CareerProfile. Profile merging is deferred to Phase 5B.
    """
    token = await register_and_login(client, "deferred.merge@example.com", "Deferred Candidate")
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Fetch initial profile
    profile_before = (await client.get("/api/v1/profile", headers=headers)).json()
    user_id = uuid.UUID(profile_before["user_id"])
    initial_exp = profile_before["years_of_experience"]
    initial_education = profile_before["education"]

    # 2. Upload and parse resume
    pdf_bytes = make_valid_pdf_bytes("Deferred Candidate")
    files = {"file": ("deferred_resume.pdf", pdf_bytes, "application/pdf")}

    upload_res = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert upload_res.status_code == 201
    assert upload_res.json()["status"] == "parsed"

    # 3. Fetch profile after parsing
    profile_after = (await client.get("/api/v1/profile", headers=headers)).json()

    # Invariants: CareerProfile must be completely untouched!
    assert profile_after["years_of_experience"] == initial_exp
    assert len(profile_after["education"]) == len(initial_education)

    # Database level check on CareerProfile
    db_profile = await test_session.scalar(select(CareerProfile).where(CareerProfile.user_id == user_id))
    assert db_profile is not None
    assert float(db_profile.years_of_experience) == initial_exp


# ─── 4. Document Retrieval, Download & Deletion Tests ─────────────────


@pytest.mark.asyncio
async def test_list_and_get_resume_metadata(client: AsyncClient):
    """Listing resumes returns all uploaded items; GET /{id} returns detailed metadata."""
    token = await register_and_login(client, "list.resumes@example.com", "List User")
    headers = {"Authorization": f"Bearer {token}"}

    pdf1 = make_valid_pdf_bytes("Resume One")
    pdf2 = make_valid_pdf_bytes("Resume Two")

    await client.post("/api/v1/resume/upload", headers=headers, files={"file": ("r1.pdf", pdf1, "application/pdf")})
    await client.post("/api/v1/resume/upload", headers=headers, files={"file": ("r2.pdf", pdf2, "application/pdf")})

    list_res = await client.get("/api/v1/resume", headers=headers)
    assert list_res.status_code == 200
    data = list_res.json()
    assert data["total"] == 2
    assert len(data["resumes"]) == 2

    # Get single resume
    r1_id = data["resumes"][0]["id"]
    detail_res = await client.get(f"/api/v1/resume/{r1_id}", headers=headers)
    assert detail_res.status_code == 200
    assert detail_res.json()["id"] == r1_id


@pytest.mark.asyncio
async def test_download_original_resume_document(client: AsyncClient):
    """Downloading a resume returns the exact stored document bytes and filename."""
    token = await register_and_login(client, "download.resume@example.com", "Download User")
    headers = {"Authorization": f"Bearer {token}"}

    original_bytes = make_valid_pdf_bytes("Download User")
    files = {"file": ("my_downloadable_resume.pdf", original_bytes, "application/pdf")}

    upload_res = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["resume_id"]

    download_res = await client.get(f"/api/v1/resume/{resume_id}/download", headers=headers)
    assert download_res.status_code == 200
    assert download_res.content == original_bytes
    assert "attachment" in download_res.headers["Content-Disposition"]
    assert "my_downloadable_resume.pdf" in download_res.headers["Content-Disposition"]


@pytest.mark.asyncio
async def test_delete_resume_cleans_db_and_storage(
    client: AsyncClient,
    test_session: AsyncSession,
):
    """Deleting a resume removes it from the database and deletes the file from storage."""
    token = await register_and_login(client, "delete.resume@example.com", "Delete User")
    headers = {"Authorization": f"Bearer {token}"}

    pdf_bytes = make_valid_pdf_bytes("Delete User")
    files = {"file": ("to_delete.pdf", pdf_bytes, "application/pdf")}

    upload_res = await client.post("/api/v1/resume/upload", headers=headers, files=files)
    assert upload_res.status_code == 201
    resume_id = uuid.UUID(upload_res.json()["resume_id"])

    # Verify stored in DB
    resume = await test_session.scalar(select(Resume).where(Resume.id == resume_id))
    assert resume is not None
    storage_path = resume.storage_path
    assert await storage_service.file_exists(storage_path) is True

    # Delete via API
    del_res = await client.delete(f"/api/v1/resume/{resume_id}", headers=headers)
    assert del_res.status_code == 204

    # Verify deleted from DB
    resume_after = await test_session.scalar(select(Resume).where(Resume.id == resume_id))
    assert resume_after is None

    # Verify deleted from filesystem
    assert await storage_service.file_exists(storage_path) is False


# ─── 5. Security & Multi-Tenant Authorization Tests ───────────────────


@pytest.mark.asyncio
async def test_cross_user_isolation_forbidden(client: AsyncClient):
    """Candidate B cannot view, download, parsed-draft, retry, or delete Candidate A's resume."""
    token_a = await register_and_login(client, "user.alpha@example.com", "Candidate Alpha")
    token_b = await register_and_login(client, "user.bravo@example.com", "Candidate Bravo")

    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}

    pdf_bytes = make_valid_pdf_bytes("Candidate Alpha")
    upload_res = await client.post("/api/v1/resume/upload", headers=headers_a, files={"file": ("alpha.pdf", pdf_bytes, "application/pdf")})
    assert upload_res.status_code == 201
    resume_a_id = upload_res.json()["resume_id"]

    # User B attempts to access User A's resume
    assert (await client.get(f"/api/v1/resume/{resume_a_id}", headers=headers_b)).status_code == 403
    assert (await client.get(f"/api/v1/resume/{resume_a_id}/parsed", headers=headers_b)).status_code == 403
    assert (await client.get(f"/api/v1/resume/{resume_a_id}/download", headers=headers_b)).status_code == 403
    assert (await client.post(f"/api/v1/resume/{resume_a_id}/retry", headers=headers_b)).status_code == 403
    assert (await client.delete(f"/api/v1/resume/{resume_a_id}", headers=headers_b)).status_code == 403
    assert (await client.post(f"/api/v1/resume/{resume_a_id}/confirm", headers=headers_b, json={"education": []})).status_code == 403


@pytest.mark.asyncio
async def test_unauthenticated_requests_rejected(client: AsyncClient):
    """Unauthenticated calls to all resume endpoints return HTTP 401 Unauthorized."""
    dummy_id = uuid.uuid4()
    assert (await client.post("/api/v1/resume/upload")).status_code == 401
    assert (await client.get("/api/v1/resume")).status_code == 401
    assert (await client.get(f"/api/v1/resume/{dummy_id}")).status_code == 401
    assert (await client.get(f"/api/v1/resume/{dummy_id}/parsed")).status_code == 401
    assert (await client.post(f"/api/v1/resume/{dummy_id}/retry")).status_code == 401
    assert (await client.get(f"/api/v1/resume/{dummy_id}/download")).status_code == 401
    assert (await client.delete(f"/api/v1/resume/{dummy_id}")).status_code == 401
    assert (await client.post(f"/api/v1/resume/{dummy_id}/confirm", json={"education": []})).status_code == 401


# ─── 6. Phase 5B: Explicit Approval & Profile Merging Tests ─────────


@pytest.mark.asyncio
async def test_confirm_resume_merges_data_into_career_profile(client: AsyncClient, test_session: AsyncSession):
    """
    Candidate explicitly approves and edits extracted resume draft:
    CareerProfile, UserProfile, and UserSkills are updated in PostgreSQL atomically.
    Extracted skills are recorded as unverified with extracted origin.
    """
    token = await register_and_login(client, "confirm.candidate@example.com", "Confirm Candidate")
    headers = {"Authorization": f"Bearer {token}"}

    # Seed canonical skills into test session if not already present
    for s_id, s_name in [("python", "Python"), ("postgresql", "PostgreSQL")]:
        sk = await test_session.scalar(select(Skill).where(Skill.id == s_id))
        if not sk:
            test_session.add(
                Skill(
                    id=s_id,
                    name=s_name,
                    category="language" if s_id == "python" else "database",
                    description=f"{s_name} programming/database competency",
                    synonyms=[s_id],
                )
            )
    await test_session.commit()

    # 1. Upload valid resume
    pdf_bytes = make_valid_pdf_bytes("Confirm Candidate")
    upload_res = await client.post(
        "/api/v1/resume/upload",
        headers=headers,
        files={"file": ("profile_resume.pdf", pdf_bytes, "application/pdf")},
    )
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["resume_id"]

    # 2. Get user id from /auth/me
    me_res = await client.get("/api/v1/auth/me", headers=headers)
    assert me_res.status_code == 200
    user_id = uuid.UUID(me_res.json()["id"])

    # 3. Submit explicit confirmation payload
    confirm_payload = {
        "personal_info": {
            "headline": "Lead Backend Architect",
            "location": "San Francisco, CA",
            "linkedin_url": "https://linkedin.com/in/confirmcandidate",
            "github_url": "https://github.com/confirmcandidate",
            "bio": "Passionate systems engineer.",
        },
        "education": [
            {
                "institution_name": "Stanford University",
                "degree": "B.S. Computer Science",
                "field_of_study": "Software Systems",
                "start_date": "2018-09-01",
                "end_date": "2022-06-15",
                "is_current": False,
                "gpa": 3.85,
            }
        ],
        "experience": [
            {
                "company_name": "TechCorp Global",
                "role_title": "Senior Backend Engineer",
                "location": "Remote",
                "start_date": "2022-07-01",
                "end_date": None,
                "is_current": True,
                "bullet_points": [
                    "Engineered distributed event pipelines processing 1M events/sec",
                    "Reduced latency by 45% using Redis caching",
                ],
            }
        ],
        "projects": [
            {
                "title": "Cloud Stream Gateway",
                "description": "High-throughput API gateway built with FastAPI and asyncpg",
                "technologies": ["Python", "FastAPI", "PostgreSQL", "Docker"],
                "github_url": "https://github.com/confirmcandidate/stream-gateway",
                "bullet_points": ["Achieved 99.99% uptime in stress testing"],
            }
        ],
        "certifications": [
            {
                "name": "AWS Certified Solutions Architect",
                "issuing_organization": "Amazon Web Services",
                "issue_date": "2023-01-15",
                "credential_id": "AWS-PSA-9921",
            }
        ],
        "confirmed_skills": ["python", "postgresql", "fastapi"],
        "make_primary": True,
    }

    confirm_res = await client.post(
        f"/api/v1/resume/{resume_id}/confirm",
        headers=headers,
        json=confirm_payload,
    )
    assert confirm_res.status_code == 200
    confirm_data = confirm_res.json()
    assert confirm_data["status"] == "confirmed"
    assert confirm_data["merged_summary"]["education_added"] == 1
    assert confirm_data["merged_summary"]["experience_added"] == 1
    assert confirm_data["merged_summary"]["projects_added"] == 1
    assert confirm_data["merged_summary"]["certifications_added"] == 1
    assert confirm_data["merged_summary"]["skills_added"] >= 2
    assert confirm_data["completeness_score"] >= 40.0

    # 4. Verify in DB directly
    career_profile = await test_session.scalar(
        select(CareerProfile).where(CareerProfile.user_id == user_id)
    )
    assert career_profile is not None

    edu_records = (await test_session.scalars(
        select(Education).where(Education.career_profile_id == career_profile.id)
    )).all()
    assert len(edu_records) == 1
    assert edu_records[0].institution_name == "Stanford University"
    assert edu_records[0].degree == "B.S. Computer Science"

    exp_records = (await test_session.scalars(
        select(Experience).where(Experience.career_profile_id == career_profile.id)
    )).all()
    assert len(exp_records) == 1
    assert exp_records[0].company_name == "TechCorp Global"

    proj_records = (await test_session.scalars(
        select(Project).where(Project.career_profile_id == career_profile.id)
    )).all()
    assert len(proj_records) == 1
    assert proj_records[0].title == "Cloud Stream Gateway"

    cert_records = (await test_session.scalars(
        select(Certification).where(Certification.career_profile_id == career_profile.id)
    )).all()
    assert len(cert_records) == 1
    assert cert_records[0].name == "AWS Certified Solutions Architect"

    # Verify user skills
    skills_records = (await test_session.scalars(
        select(UserSkill).where(UserSkill.career_profile_id == career_profile.id)
    )).all()
    assert len(skills_records) >= 2
    for us in skills_records:
        assert us.primary_source == "extracted"
        assert us.proficiency_tier == "unverified"
        assert us.is_verified is False


@pytest.mark.asyncio
async def test_confirm_resume_idempotent_duplicate_prevention(client: AsyncClient, test_session: AsyncSession):
    """Confirming identical payload again avoids creating duplicate records."""
    token = await register_and_login(client, "idempotent.candidate@example.com", "Idempotent Candidate")
    headers = {"Authorization": f"Bearer {token}"}

    sk = await test_session.scalar(select(Skill).where(Skill.id == "python"))
    if not sk:
        test_session.add(
            Skill(
                id="python",
                name="Python",
                category="language",
                description="Python programming competency",
                synonyms=["python"],
            )
        )
        await test_session.commit()

    pdf_bytes = make_valid_pdf_bytes("Idempotent Candidate")
    upload_res = await client.post(
        "/api/v1/resume/upload",
        headers=headers,
        files={"file": ("idempotent.pdf", pdf_bytes, "application/pdf")},
    )
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["resume_id"]

    confirm_payload = {
        "education": [
            {
                "institution_name": "MIT",
                "degree": "M.S.",
                "field_of_study": "AI",
                "start_date": "2020",
                "is_current": False,
            }
        ],
        "confirmed_skills": ["python"],
    }

    # First confirmation
    res1 = await client.post(f"/api/v1/resume/{resume_id}/confirm", headers=headers, json=confirm_payload)
    assert res1.status_code == 200
    assert res1.json()["merged_summary"]["education_added"] == 1

    # Second confirmation with exact same data
    res2 = await client.post(f"/api/v1/resume/{resume_id}/confirm", headers=headers, json=confirm_payload)
    assert res2.status_code == 200
    # Duplicate detected, 0 added
    assert res2.json()["merged_summary"]["education_added"] == 0
    assert res2.json()["merged_summary"]["skills_added"] == 0


@pytest.mark.asyncio
async def test_confirm_unparsed_resume_rejected(client: AsyncClient):
    """Attempting to confirm a resume that failed parsing returns HTTP 409 Conflict."""
    token = await register_and_login(client, "unparsed.candidate@example.com", "Unparsed Candidate")
    headers = {"Authorization": f"Bearer {token}"}

    # Upload corrupt PDF that will fail parsing
    corrupt_bytes = b"%PDF-1.4\nBROKEN_DOCUMENT_NOT_PARSEABLE"
    upload_res = await client.post(
        "/api/v1/resume/upload",
        headers=headers,
        files={"file": ("broken.pdf", corrupt_bytes, "application/pdf")},
    )
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["resume_id"]

    # Confirm should return 409
    confirm_res = await client.post(
        f"/api/v1/resume/{resume_id}/confirm",
        headers=headers,
        json={"education": []},
    )
    assert confirm_res.status_code == 409
